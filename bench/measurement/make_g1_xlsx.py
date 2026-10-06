#!/usr/bin/env python3
"""Build the human-ready G1 labelling workbook from the AI-drafted packet.

Turns a packet CSV (independent-model first-pass labels) into an ergonomic .xlsx a
co-author can adjudicate in about an hour:

  * dropdown-validated label cells,
  * the AI draft and its rationale parked on the right as a reference,
  * a BLIND subset floated to the top so the honest (un-anchored) kappa can be
    computed,
  * and, for every row, the SOURCE FILE the value came from plus a clickable
    GitHub permalink, so the annotator opens the exact config instead of searching
    the repo.

Fill it, save the Label sheet back to CSV, and grade with grade_dual.py. See
LABELING.md.

  python3 bench/measurement/make_g1_xlsx.py \
      bench/measurement/out-val/labeling-packet-sample-ai.csv \
      bench/measurement/out-val/g1-packet.xlsx \
      [--blind 40] [--corpus bench/measurement/corpus] [--sources bench/measurement/sources.tsv]
"""
import argparse
import csv
import glob
import hashlib
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

LABELS = ["correct", "correct-extra", "discovery-miss", "parser-artifact",
          "no-consumer", "spurious", "wrong-attribution"]
BLIND_SEED = "keyway-g1-blind"
MAX_FILES = 3  # cap the source-file list per row

# Label sheet columns, in review order: identity + evidence (incl. the source file
# and a GitHub link) first, the two human columns next, then the AI reference
# pushed to the far right so it can be ignored during the blind pass.
COLS = [
    ("row_id", 6), ("blind_first", 11), ("repo", 24), ("field", 9),
    ("value", 26), ("source_file", 44), ("github", 9), ("side", 14),
    ("in_nonauth_context", 10), ("has_jwt_consumer", 10),
    ("human_label_blind", 18), ("human_label", 18), ("notes", 20),
    ("ai_label", 16), ("ai_confidence", 8), ("ai_rationale", 68),
    ("suggested_label", 16),
]
HUMAN_COLS = {"human_label_blind", "human_label"}


def blind_set(rows, n):
    ranked = sorted(
        range(len(rows)),
        key=lambda i: hashlib.sha256(
            f"{BLIND_SEED}|{rows[i]['repo']}|{rows[i]['field']}|{rows[i]['value']}"
            .encode()).hexdigest())
    return set(ranked[:min(n, len(rows))])


def _sanitize(path):
    return re.sub(r"[^A-Za-z0-9._-]", "_", path)


def load_url_map(sources_tsv):
    """corpus-basename -> GitHub blob permalink, reconstructed from the manifest."""
    url = {}
    if not os.path.exists(sources_tsv):
        return url
    with open(sources_tsv) as fh:
        for r in csv.DictReader(fh, delimiter="\t"):
            e = f'{r["repo"].replace("/", "_")}__{_sanitize(r["path"])}'
            for key in (e, e + ".yaml", e + ".yml"):
                url.setdefault(key, r["url"])
    return url


def repo_files(corpus_dir, repo):
    """[(basename, text)] for a repo prefix, read once and cached by the caller."""
    out = []
    for p in sorted(glob.glob(os.path.join(corpus_dir, f"{repo}__*"))):
        try:
            out.append((os.path.basename(p), open(p, errors="ignore").read()))
        except OSError:
            pass
    return out


def locate(files, field, value):
    """Basenames of the repo's files that contain this (field, value)."""
    hits = []
    for base, text in files:
        if field == "claims":
            if (f"claims[{value}" in text or f"claims.{value}" in text
                    or (value in text and "claims" in text)):
                hits.append(base)
        elif value and value in text:
            hits.append(base)
    return hits


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("inp", nargs="?", default="bench/measurement/out-val/labeling-packet-sample-ai.csv")
    ap.add_argument("out", nargs="?", default="bench/measurement/out-val/g1-packet.xlsx")
    ap.add_argument("--blind", type=int, default=40)
    ap.add_argument("--corpus", default="bench/measurement/corpus")
    ap.add_argument("--sources", default="bench/measurement/sources.tsv")
    args = ap.parse_args()

    rows = list(csv.DictReader(open(args.inp)))

    # Attach the source file(s) and GitHub link per row.
    url_map = load_url_map(args.sources)
    cache, located = {}, 0
    for r in rows:
        repo = r["repo"]
        if repo not in cache:
            cache[repo] = repo_files(args.corpus, repo)
        bases = locate(cache[repo], r["field"], r["value"])
        if bases:
            located += 1
        r["source_file"] = "; ".join(bases[:MAX_FILES]) + (
            f"  (+{len(bases) - MAX_FILES} more)" if len(bases) > MAX_FILES else "")
        r["_url"] = next((url_map.get(b) for b in bases if url_map.get(b)), "")
        r["github"] = "open ↗" if r["_url"] else ""

    blind = blind_set(rows, args.blind)
    order = [i for i in range(len(rows)) if i in blind] + \
            [i for i in range(len(rows)) if i not in blind]

    wb = Workbook()
    _instructions(wb.active, len(rows), len(blind))
    _label_sheet(wb.create_sheet("Label"), rows, order, blind)
    wb.save(args.out)
    print(f"wrote {args.out}: {len(rows)} rows ({len(blind)} blind-first); "
          f"source file located for {located}/{len(rows)}")


def _instructions(ws, n, nblind):
    ws.title = "START HERE"
    ws.sheet_view.showGridLines = False
    big = Font(bold=True, size=14)
    hd = Font(bold=True, size=11)
    wrap = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100

    def row(r, a, b="", font=None, w=False):
        ws[f"A{r}"] = a
        ws[f"B{r}"] = b
        if font:
            ws[f"A{r}"].font = font
        if w:
            ws[f"B{r}"].alignment = wrap

    row(1, "Keyway G1", "Human validation packet (gold-standard discovery precision / recall)", big)
    row(3, "What this is",
        f"A stratified random sample of {n} (repo, field, value) rows across the "
        f"corpus. You adjudicate each with one label from the taxonomy below; the "
        f"labels give population precision and recall with confidence intervals. The "
        f"'source_file' column names the config the value came from, and 'github' "
        f"links straight to it, so you do not have to search the repo.", hd, w=True)
    row(6, "Do this, in order", "", hd)
    row(7, "  1. BLIND pass",
        f"The first {nblind} rows are marked BLIND. Fill ONLY their 'human_label_blind' "
        f"cell, using the value, source file, side and flags and your own judgement. "
        f"Do NOT scroll right to the ai_label / ai_rationale columns yet. This gives "
        f"an un-anchored agreement number.", w=True)
    row(10, "  2. Full pass",
        "Now fill 'human_label' on EVERY row (including the blind ones). Here you may "
        "use the ai_label and ai_rationale columns (far right) as a reference and "
        "correct them. Open 'github' (or the source_file locally) to check the config. "
        "Both label cells are dropdowns.", w=True)
    row(14, "  3. Grade",
        "File > Save As > CSV (the 'Label' sheet) to e.g. labeling-packet-human.csv, then:", w=True)
    row(15, "", "python3 bench/measurement/grade_dual.py bench/measurement/out-val/labeling-packet-human.csv")
    row(16, "", "It reports v1 (AI) vs v2 (human) precision/recall, Cohen's kappa, and the "
        "blind-subset kappa (the honest, un-anchored one).", w=True)

    row(18, "Label taxonomy", "fill with exactly one", hd)
    tax = [
        ("correct", "real auth config, correctly captured/attributed -> TP"),
        ("correct-extra", "captured-only, but the value IS real (parser missed it) -> TP"),
        ("discovery-miss", "real auth value Keyway FAILED to capture -> hurts recall (FN)"),
        ("parser-artifact", "parser scraped a non-auth value (ConfigMap/annotation/comment) -> excluded"),
        ("no-consumer", "real value but its RequestAuthentication isn't in the corpus -> excluded"),
        ("spurious", "Keyway captured a value that is NOT real auth config -> hurts precision (FP)"),
        ("wrong-attribution", "real value attached to the WRONG consumer -> hurts precision (FP)"),
    ]
    for i, (lab, desc) in enumerate(tax):
        r = 19 + i
        ws[f"A{r}"] = lab
        ws[f"A{r}"].font = Font(bold=True)
        ws[f"B{r}"] = desc
    r = 19 + len(tax)
    row(r + 1, "Metrics", "recall = TP / (TP + discovery-miss);  "
        "precision = TP / (TP + spurious + wrong-attribution). "
        "parser-artifact and no-consumer are excluded from both.", hd, w=True)
    row(r + 3, "Two annotators",
        "For inter-rater agreement, a second co-author fills their own copy; compare "
        "the two 'human_label' columns. Strongest on the blind subset.", w=True)


def _label_sheet(ws, rows, order, blind):
    names = [c[0] for c in COLS]
    hb_col = get_column_letter(1 + names.index("human_label_blind"))
    ws.freeze_panes = f"{hb_col}2"  # keep header + identity/evidence/source visible
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="334155")
    blind_fill = PatternFill("solid", fgColor="FEF3C7")
    link_font = Font(color="2563EB", underline="single")
    wrap = Alignment(wrap_text=True, vertical="top")

    for c, (name, width) in enumerate(COLS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = head_font
        cell.fill = head_fill
        ws.column_dimensions[get_column_letter(c)].width = width

    for out_r, i in enumerate(order, start=2):
        src = rows[i]
        is_blind = i in blind
        rec = {
            "row_id": out_r - 1,
            "blind_first": "BLIND" if is_blind else "",
            "repo": src.get("repo", ""),
            "field": src.get("field", ""),
            "value": src.get("value", ""),
            "source_file": src.get("source_file", ""),
            "github": src.get("github", ""),
            "side": src.get("side", ""),
            "in_nonauth_context": src.get("in_nonauth_context", ""),
            "has_jwt_consumer": src.get("has_jwt_consumer", ""),
            "human_label_blind": "",
            "human_label": "",
            "notes": "",
            "ai_label": src.get("ai_label", ""),
            "ai_confidence": src.get("ai_confidence", ""),
            "ai_rationale": src.get("ai_rationale", ""),
            "suggested_label": src.get("suggested_label", ""),
        }
        for c, (name, _) in enumerate(COLS, start=1):
            cell = ws.cell(row=out_r, column=c, value=rec[name])
            if name in ("value", "source_file", "ai_rationale"):
                cell.alignment = wrap
            if name == "github" and src.get("_url"):
                cell.hyperlink = src["_url"]
                cell.font = link_font
            if name == "blind_first" and is_blind:
                cell.fill = blind_fill
                cell.font = Font(bold=True, color="92400E")

    dv = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    last = len(order) + 1
    for name in HUMAN_COLS:
        col = get_column_letter(1 + names.index(name))
        dv.add(f"{col}2:{col}{last}")


if __name__ == "__main__":
    main()
